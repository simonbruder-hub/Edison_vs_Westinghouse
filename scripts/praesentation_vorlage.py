"""Füllt die Slidesgo-Vorlage «Reconstruction Era and the Gilded Age» mit dem Referat.

Eingabe: die bereits strukturierte Vorlage (Folien ausgewählt, Zeitstrahl- und Textfolie dupliziert,
Reihenfolge 1, 4, 17, 11, 25, 25b, 9, 29, 19, 30, 6, 35, 6b).
Aufruf:  python3 scripts/praesentation_vorlage.py struct.pptx abgabe/Praesentation_Stromkrieg_Bruder.pptx
Zitatnummern = Nummern im Handout (assets/zitierreihenfolge.json).
"""
import copy
import json
import re
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Pt

ROOT = Path(__file__).resolve().parent.parent
ORDER = json.loads((ROOT / 'assets' / 'zitierreihenfolge.json').read_text())
A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'


def n(*keys):
    nums = sorted(ORDER.index(k) + 1 for k in keys)
    parts, i = [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        parts += [f'[{nums[i]}]–[{nums[j]}]'] if j - i >= 2 else [f'[{x}]' for x in nums[i:j + 1]]
        i = j + 1
    return ', '.join(parts)


def set_paras(tf, texts, size=None):
    """Ersetzt den Text, behält Absatz- und Zeichenformat der Vorlage (Absatz i nutzt Vorlagenabsatz i)."""
    body = tf._txBody
    old = body.findall(A + 'p')
    for k, text in enumerate(texts):
        tpl = old[min(k, len(old) - 1)]
        p = copy.deepcopy(tpl)
        runs = p.findall(A + 'r')
        for el in list(p):
            if el.tag in (A + 'r', A + 'br', A + 'fld'):
                p.remove(el)
        r = copy.deepcopy(runs[0]) if runs else etree.SubElement(p, A + 'r')
        rpr = r.find(A + 'rPr')
        if rpr is not None:
            for h in rpr.findall(A + 'hlinkClick'):
                rpr.remove(h)
            if size:
                rpr.set('sz', str(int(size * 100)))
        r.find(A + 't').text = text
        end = p.find(A + 'endParaRPr')
        if end is not None:
            end.addprevious(r)
        else:
            p.append(r)
        body.append(p)
    for p in old:
        body.remove(p)


def shapes_by_id(slide):
    out = {}

    def walk(shs):
        for sh in shs:
            out[sh.shape_id] = sh
            if sh.shape_type == 6:
                walk(sh.shapes)
    walk(slide.shapes)
    return out


def fill(slide, mapping, notes):
    sh = shapes_by_id(slide)
    for sid, val in mapping.items():
        size = None
        if isinstance(val, tuple):
            val, size = val
        set_paras(sh[sid].text_frame, val if isinstance(val, list) else [val], size)
    slide.notes_slide.notes_text_frame.text = notes


def fill_table(table, rows, size=None):
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            set_paras(table.cell(r, c).text_frame, [text], size)


src, out = sys.argv[1], sys.argv[2]
prs = Presentation(src)
S = prs.slides

# 1 Titel
fill(S[0], {895: 'FHNW', 890: ['Der Stromkrieg', '- 1886–1892 -'],
            891: 'Edison vs. Westinghouse · Simon Bruder'},
     'Titelfolie – kurz begrüssen, Thema nennen, direkt zum Einstieg weiter.')

# 2 Einstieg
fill(S[1], {1864: ('- Auburn, 1890 -', 54),
            1865: '6. August: erste Hinrichtung mit Strom – mit einem Westinghouse-Generator'},
     f"""[0:00–0:45] Einstieg
Auburn, New York, 6. August 1890: William Kemmler wird als erster Mensch auf dem elektrischen Stuhl hingerichtet. Der Generator stammt von Westinghouse – beschafft hatte ihn aber das Lager um Edison, um Wechselstrom als «Todesstrom» abzustempeln.
Leitfrage: Warum setzte sich in den USA trotzdem der Wechselstrom durch – und was zeigt der Konflikt über den Aufstieg der USA zur führenden Industriemacht?
These: Entschieden haben nicht Propaganda oder Genies, sondern Kosten, Kapital und Netzlogik.
Quellen: {n('moran', 'essig')}.""")

# 3 Technik (Tabelle)
s = S[2]
fill(s, {3480: 'Warum die Spannung entscheidet'}, f"""[0:45–1:45] Technik
Kernidee: Strom lässt sich nur mit hoher Spannung billig über weite Strecken schicken.
Der Verlust in der Leitung wächst mit dem Quadrat der Stromstärke. Zehnfache Spannung heisst bei gleicher Leistung ein Zehntel des Stroms – und nur noch ein Hundertstel des Verlusts.
Um 1890 liess sich nur Wechselstrom mit Transformatoren hoch- und heruntersetzen. Edisons Gleichstrom brauchte viele kleine Kraftwerke mitten in der Stadt und viel Kupfer.
Quellen: {n('pearl', 'edpCW', 'jonnes')}.""")
tbl = shapes_by_id(s)[3481].table
fill_table(tbl, [
    ['', 'Gleichstrom (Edison)', 'Wechselstrom (Westinghouse)'],
    ['Spannung', 'rund 110 V, nicht transformierbar', 'Transformator: hoch im Netz, tief beim Kunden'],
    ['Verlust (I² · R)', 'hoch: viel Strom, dicke Kupferleitungen', '10 × Spannung = 1/100 Verlust'],
    ['Reichweite', f'ein Stadtquartier (1882: 0,65 km²) {n("pearl")}', f'Niagara–Buffalo 1896: über 30 km {n("jonnes")}'],
    ['Schwachpunkt', 'viele kleine Kraftwerke, viel Kupfer', 'Hochspannung gefährlich; Motor erst 1888'],
], size=12)
for row in tbl.rows:
    row.height = int(0.5 * 914400)

# 4 Akteure
fill(S[3], {
    2429: 'Firmen und Investoren statt Genies',
    2434: ('Thomas A. Edison', 15), 2430: 'Erfinder-Unternehmer; verteidigt Gleichstrom',
    2436: ('G. Westinghouse', 15), 2431: 'Industrieller; ab 1886 Wechselstromnetze',
    2435: ('Nikola Tesla', 15), 2432: f'Motorpatente 1888 für die Fabrik {n("carlson")}',
    2437: ('Harold P. Brown', 15), 2433: f'Anti-Wechselstrom-Kampagne {n("hughes58")}',
}, f"""[1:45–3:00] Akteure
Im Kern stritten zwei Aktiengesellschaften mit ihren Investoren, nicht bloss zwei Erfinder {n('edpCW')}.
Edison (1847–1931) hatte viel Geld und Ruf in Gleichstrom-Kraftwerke gesteckt. Seine Warnungen vor Hochspannung waren nicht nur Taktik – er hielt sie auch für gefährlich.
Westinghouse (1846–1914), bekannt durch die Druckluftbremse, war Industrieller: Er kaufte die nötigen Patente zu, darunter 1888 Teslas Motorpatente.
Tesla (1856–1943), 1884 kurz bei Edison, lieferte mit dem Wechselstrommotor das fehlende Stück für Fabriken.
Harold P. Brown war das öffentliche Gesicht der Kampagne gegen Wechselstrom; Edisons Labor half im Hintergrund.""")

# 5 Verlauf I (Zeitstrahl)
fill(S[4], {
    4358: '1882–1890: Wettbewerb und Propaganda',
    4368: '1882', 4361: ('Pearl Street', 13), 4362: f'Edisons Gleichstrom für 400 Lampen {n("pearl")}',
    4370: '1886', 4365: ('Great Barrington', 13), 4366: f'erstes Wechselstromnetz der USA {n("jonnes")}',
    4367: '1888', 4359: ('«A Warning»', 13), 4360: f'Brown tötet öffentlich Hunde {n("hughes58")}',
    4369: '1889/90', 4363: ('Elektrischer Stuhl', 13), 4364: f'Kemmler-Hinrichtung als Debakel {n("moran", "essig")}',
}, f"""[3:00–4:15] Verlauf, Teil 1
Phase 1, offener Wettbewerb: 1882 startet Edisons Pearl Street Station, 1886 läuft in Great Barrington das erste Wechselstromnetz der USA mit Transformatoren; im selben Jahr gründet Westinghouse seine Elektrofirma.
Phase 2, Propagandakrieg: 1888 warnt Edisons Firma mit der Broschüre «A Warning», Harold Brown tötet öffentlich Hunde mit Wechselstrom.
1889 führt New York den elektrischen Stuhl ein – bewusst mit einem Westinghouse-Generator. Edison und Westinghouse streiten sich zudem in der North American Review {n('edison89', 'west89')}.
1890 die Hinrichtung Kemmlers: kein «sauberer» Tod, sondern ein Debakel – zurück zum Einstieg.""")

# 6 Verlauf II (Zeitstrahl)
fill(S[5], {
    4358: '1891–1896: Markt und Kapital entscheiden',
    4368: '1891', 4361: ('Lizenzverzicht', 13), 4362: f'Tesla verzichtet auf Gebühren {n("carlson")}',
    4370: '1892', 4365: ('General Electric', 13), 4366: f'Fusion unter J. P. Morgan {n("edpCW")}',
    4367: '1893', 4359: ('Chicago', 13), 4360: f'Westinghouse beleuchtet die Weltausstellung {n("jonnes")}',
    4369: '1896', 4363: ('Niagara–Buffalo', 13), 4364: f'Fernstrom und Patentabkommen {n("jonnes", "carlson")}',
}, """[4:15–5:30] Verlauf, Teil 2
Phase 3: Jetzt entscheiden Markt und Kapital.
1891 ist Westinghouse finanziell unter Druck; Tesla verzichtet auf seine Lizenzgebühren.
Wendepunkt 1892: Der Bankier J. P. Morgan fusioniert Edisons Firma mit Thomson-Houston zu General Electric. Edison verliert die Kontrolle – und GE setzt ebenfalls auf Wechselstrom. Damit ist der Systemstreit faktisch entschieden.
1893 beleuchtet Westinghouse die Weltausstellung in Chicago, deutlich günstiger als GE. 1896 fliesst Strom von den Niagarafällen nach Buffalo; im selben Jahr teilen GE und Westinghouse ihre Patente – aus Rivalen wird ein Duopol.""")

# 7 Warum
fill(S[6], {
    2378: 'Kosten und Kapital schlugen Propaganda',
    2382: 'Kosten', 2379: f'Wechselstromnetze waren billiger; neue Versorger und ab 1892 auch GE setzten darauf {n("edpCW")}',
    2383: 'Propaganda scheitert', 2380: f'Wechselstrom liess sich öffentlich nicht mit dem Tod verbinden {n("cole")}',
    2384: 'Schweizer Bezug', 2381: f'1891: Drehstrom über 175 km (MFO Oerlikon, AEG); Brown gründet BBC in Baden {n("steig")}',
}, f"""[5:30–6:30] Warum Wechselstrom gewann
Erstens Kosten: Wenige grosse Kraftwerke und dünnere Leitungen waren billiger als viele kleine Gleichstromzentralen.
Zweitens: Die Angstkampagne scheiterte. Cole und Chandler (2019) beschreiben sie als «kompetitives Impression Management» – Edisons Lager wollte das Bild des Konkurrenzprodukts beschädigen, schaffte es aber nicht, Wechselstrom in der Öffentlichkeit dauerhaft mit dem Tod zu verknüpfen.
Drittens der Schweizer Bezug: 1891 übertragen AEG und die Maschinenfabrik Oerlikon Drehstrom über 175 km von Lauffen nach Frankfurt. Der Oerlikon-Ingenieur Charles E. L. Brown gründet im selben Jahr in Baden die BBC – heute ABB.
Kein K.-o.-Sieg: Umformer verbanden alte und neue Netze; in Manhattan lief Gleichstrom bis 2007 {n('davidBunn', 'millard', 'lee')}.
Deutung: Paul David nennt das Pfadabhängigkeit – ein Standard setzt sich in einem kurzen, offenen Zeitfenster durch und verfestigt sich dann {n('david92')}.""")

# 8 Grafik Weltindustrie
s = S[7]
fill(s, {
    4456: 'Die USA überholen Grossbritannien',
    4457: (f'Anteil an der Weltindustrieproduktion in %, Stichjahre (Zeitachse nicht massstäblich). Daten: Bairoch {n("bairoch")}', 9),
    4458: 'USA', 4459: ('1913: 32 % der Weltindustrie', 12),
    4460: ('Grossbritannien', 13), 4461: ('1880: 22,9 %, 1913: 13,6 %', 12),
}, f"""[6:30–7:30] Bedeutung für die USA (1)
Bezug zum Modul «Weltmacht USA». Die Daten stammen vom Wirtschaftshistoriker Paul Bairoch.
1880 lag Grossbritannien klar vorne. Zwischen 1880 und 1900 – im Jahrzehnt des Stromkriegs – ziehen die USA vorbei. 1913 stellen sie knapp ein Drittel der Weltindustrieproduktion, mehr als Grossbritannien und Deutschland zusammen (28,4 %).
Vorsicht bei der Deutung: Der Stromkrieg hat das nicht allein verursacht. Aber billiger, übertragbarer Strom wurde zu einer Basistechnologie dieses Aufstiegs {n('hughes91')}.""")
pic = shapes_by_id(s)[4464]
rid = pic._element.blipFill.find(A + 'blip').get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
s.part.related_part(rid)._blob = (ROOT / 'assets' / 'chart_weltindustrie_vorlage.png').read_bytes()
# Legendenquadrate: oben = USA (dunkles Navy), unten = Grossbritannien
sq = shapes_by_id(s)
for sid, col in ((4462, '0A1C5E'), (4463, '4E5B9B')):
    sq[sid].fill.solid()
    sq[sid].fill.fore_color.rgb = __import__('pptx.dml.color', fromlist=['RGBColor']).RGBColor.from_string(col)

# 9 Zahlen Fabrik und Konzern
fill(S[8], {
    3953: '5 % → 78 %', 3952: f'Elektromotoren an der Antriebsleistung der US-Industrie, 1899 → 1929 {n("devine")}',
    3954: '1,3 % → 3,1 %', 3955: f'Produktivitätswachstum pro Jahr, vor → nach 1919 {n("david90")}',
    3956: '1896', 3957: 'GE unter den zwölf ersten Werten des Dow Jones',
}, f"""[7:30–8:30] Bedeutung für die USA (2)
Fabriken: 1899 lieferten Elektromotoren knapp 5 Prozent der Antriebsleistung der US-Industrie, 1929 rund 78 Prozent.
Die Produktivität wuchs vor 1919 um 1,3 Prozent pro Jahr, danach um 3,1 Prozent. Paul David erklärt die Verzögerung: Erst als Fabriken um viele Elektromotoren herum neu gebaut wurden, zahlte sich die Technik aus.
Big Business: GE war 1896 einer der zwölf ersten Werte im Dow Jones; mit Westinghouse bildete es ein Duopol mit Patentpool – typisch für das Gilded Age. An den Niagarafällen zog billige Wasserkraft Aluminium- und Chemieindustrie an {n('hughes83')}.""")

# 10 Mythen (drei Tabellen: Mythos / Befund)
s = S[9]
fill(s, {4469: 'Drei Mythen – und was die Quellen sagen'}, f"""[8:30–9:30] Mythen und Quellenkritik
Erstens Topsy: Ein ABB-Beitrag behauptet, Edison habe zur Abschreckung einen Elefanten töten lassen. Die Edison Papers der Rutgers University zeigen: Das war 1903, über zehn Jahre nach dem Stromkrieg, und Edison war nicht beteiligt.
Zweitens «Tesla gegen Edison»: Filme wie «The Current War» machen daraus ein Duell der Genies; tatsächlich stritten Firmen und Investoren.
Drittens der zerrissene Vertrag: Dafür gibt es keinen Beleg – belegt ist nur der Verzicht auf Lizenzgebühren (1891) und der spätere Patentverkauf für 216 000 Dollar (Carlson).""")
sh = shapes_by_id(s)
fill_table(sh[4470].table, [[f'Mythos: Edison liess 1903 den Elefanten Topsy töten {n("abb")}'],
                            [f'Befund: über zehn Jahre nach dem Stromkrieg, ohne Edisons Beteiligung {n("topsy")}']], size=12)
fill_table(sh[4471].table, [['Mythos: Der Stromkrieg war ein Duell «Tesla gegen Edison»'],
                            [f'Befund: Gestritten haben vor allem Unternehmen und Investoren {n("edpCW")}']], size=12)
fill_table(sh[4472].table, [['Mythos: Tesla zerriss seinen Vertrag mit Westinghouse'],
                            [f'Befund: Belegt sind nur Lizenzverzicht (1891) und Patentverkauf {n("carlson")}']], size=12)

# 11 Fazit
fill(S[10], {
    2340: 'Fazit: der erste grosse Standardkrieg',
    2342: f'Gewonnen hat das System, das sich billiger skalieren liess – nicht die bessere Propaganda. Mitgewonnen hat der finanzstarke Grosskonzern: GE und Westinghouse wurden zum Muster des US-Industriekapitalismus {n("mcnichol")}.',
    2341: f'Ironie der Geschichte: Mit Stromrichtern kehrt der Gleichstrom zurück. Die erste kommerzielle Hochspannungs-Gleichstrom-Übertragung (Gotland, 1954) baute ASEA – heute wie die BBC aus Baden Teil von ABB {n("gotland")}.',
}, """[9:30–10:00] Fazit
Zurück zur Leitfrage: Wechselstrom gewann, weil er sich billiger skalieren liess und weil das grosse Kapital – J. P. Morgan und GE – am Ende darauf setzte. Die Angstkampagne hat daran nichts geändert.
Für die «Weltmacht USA»: Der Stromkrieg steht am Anfang eines Musters – Standardkriege, die mit Kapital entschieden werden, und Grosskonzerne, die danach den Markt aufteilen.
Pointe: Heute kehrt der Gleichstrom für lange Strecken zurück – mit HGÜ, zuerst 1954 auf Gotland, gebaut von ASEA, heute wie die BBC Teil von ABB.""")

# 12 Danke
s = S[11]
sh = shapes_by_id(s)
fill(s, {4515: 'Danke!', 4516: ['Fragen?', 'Simon Bruder · Modul «Weltmacht USA» (FHNW)', 'Dozent: Stephan Schwarz · 10.12.2026']}, 'Danke – Fragen?')
for sid in (4845, 4848, 4851, 4854, 4855, 4860):  # Social-Media-Icons der Vorlage entfernen
    el = sh[sid]._element
    el.getparent().remove(el)

# 13 Quellen (Auswahl)
REFS_SHORT = {
    'edpCW': 'Thomas A. Edison Papers, “The current wars,” Rutgers Univ., online.',
    'jonnes': 'J. Jonnes, Empires of Light. New York: Random House, 2003.',
    'carlson': 'W. B. Carlson, Tesla: Inventor of the Electrical Age. Princeton Univ. Press, 2013.',
    'hughes58': 'T. P. Hughes, “Harold P. Brown and the executioner’s current,” Bus. Hist. Rev., vol. 32, no. 2, 1958.',
    'cole': 'B. M. Cole and D. Chandler, “A model of competitive impression management,” Adm. Sci. Q., vol. 64, no. 4, 2019.',
    'david92': 'P. A. David, “Heroes, herds and hysteresis in technological history,” Ind. Corp. Change, vol. 1, no. 1, 1992.',
    'bairoch': 'P. Bairoch, “International industrialization levels from 1750 to 1980,” J. Eur. Econ. Hist., vol. 11, no. 2, 1982.',
    'devine': 'W. D. Devine, Jr., “From shafts to wires,” J. Econ. Hist., vol. 43, no. 2, 1983.',
    'david90': 'P. A. David, “The dynamo and the computer,” Am. Econ. Rev., vol. 80, no. 2, 1990.',
    'topsy': 'Thomas A. Edison Papers, “Myth buster: Topsy the elephant,” Rutgers Univ., online.',
}
keys = sorted(REFS_SHORT, key=lambda k: ORDER.index(k))
lines = [f'[{ORDER.index(k) + 1}] {REFS_SHORT[k]}' for k in keys]
fill(S[12], {2340: 'Quellen (Auswahl, IEEE)', 2342: (lines[:5], 11), 2341: (lines[5:], 11)},
     'Anhang – nur bei Nachfragen zeigen. Vollständiges IEEE-Verzeichnis (30 Einträge) im Handout; Nummern identisch.')

sh = shapes_by_id(S[12])
for sid in (2340,):
    pass
for sid in (2341, 2342):
    sh[sid].top, sh[sid].height = int(1.75 * 914400), int(2.9 * 914400)
    sh[sid].text_frame._txBody.bodyPr.set('anchor', 't')

prs.save(out)
print('OK', out)
